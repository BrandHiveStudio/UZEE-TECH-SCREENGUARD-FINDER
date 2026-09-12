import openpyxl
import json
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('UZEE_TECH_SUPER_D_MASTER_263_WITH_DISPLAY_SIZE_QA.xlsx', data_only=True)
ws = wb['MASTER_263']

excel_ds = {}
for r in range(2, ws.max_row + 1):
    gid = ws.cell(row=r, column=1).value
    dsize = ws.cell(row=r, column=5).value
    if gid and dsize and str(dsize).strip() != 'Unknown':
        excel_ds[str(gid).strip()] = str(dsize).strip()

print(f"Excel groups with non-Unknown display size: {len(excel_ds)}")

with open('scripts/display_size_updates.json', 'r', encoding='utf-8') as f:
    json_updates = json.load(f)

print(f"JSON updates count: {len(json_updates)}")

with open('scripts/update_display_sizes_263qa.sql', 'r', encoding='utf-8') as f:
    sql_text = f.read()

lines = [line.strip() for line in sql_text.split('\n') if line.startswith('UPDATE boxes')]
print(f"SQL update statements count: {len(lines)}")

mismatches = 0
for line in lines:
    m = re.search(r"display_size = '([^']+)' WHERE id = '([^']+)'", line)
    if m:
        ds, gid = m.group(1), m.group(2)
        expected = excel_ds.get(gid)
        if expected != ds:
            print(f"Mismatch for {gid}: SQL has {ds}, Excel has {expected}")
            mismatches += 1

if mismatches == 0 and len(lines) == 70:
    print("✅ ALL 70 SQL updates PERFECTLY MATCH the authoritative Excel workbook!")
