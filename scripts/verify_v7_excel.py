import openpyxl
import pandas as pd
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_name = 'UZEE_TECH_SUPER_D_UNASSIGNED_CLEAN_V7.xlsx'
wb = openpyxl.load_workbook(file_name, data_only=True)

print("Workbook Sheet Names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    df = pd.read_excel(file_name, sheet_name=sname)
    print(f"\n--- Sheet '{sname}' ---")
    print(f"  Rows (excl header): {ws.max_row - 1}, Cols: {ws.max_column}")
    print("  Columns:", df.columns.tolist())
    print("  First 2 rows:")
    print(df.head(2).to_string())
    print("  Freeze panes:", ws.freeze_panes)
    print("  AutoFilter ref:", ws.auto_filter.ref)

print("\n-----------------------------------------------------")
print("DATA INTEGRITY VERIFICATION:")

df_s1 = pd.read_excel(file_name, sheet_name='CLEAN UNASSIGNED GROUPS')
df_s2 = pd.read_excel(file_name, sheet_name='MODEL NORMALIZATION')
df_s3 = pd.read_excel(file_name, sheet_name='NEEDS MANUAL VERIFICATION')
df_s4 = pd.read_excel(file_name, sheet_name='SOURCE MODEL COMPARISON')
df_s5 = pd.read_excel(file_name, sheet_name='STATISTICS')

print(f"Unassigned Groups Processed: {len(df_s1)}")
print(f"Model Normalizations Logged: {len(df_s2)}")
print(f"Models Requiring Manual Verification: {len(df_s3)}")
print(f"Source Model Comparisons Logged: {len(df_s4)}")
print(f"Statistics Rows: {len(df_s5)}")

assert len(df_s1) == 163, f"Expected 163 unassigned groups, got {len(df_s1)}"
assert len(df_s4) == 719, f"Expected 719 model comparisons, got {len(df_s4)}"

print("\n✅ V7 QUALITY PASS VERIFICATION PASSED PERFECTLY!")
