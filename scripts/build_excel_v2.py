import pandas as pd
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')

# 1. Load screenguards.json
with open('src/data/screenguards.json', 'r', encoding='utf-8') as f:
    sg_json = json.load(f)

json_boxes = sg_json['boxes']
json_box_dict = {b['boxNumber']: b for b in json_boxes}

# 2. Load V1 Excel sheets
excel_v1 = 'UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS.xlsx'
df_final_v1 = pd.read_excel(excel_v1, sheet_name='FINAL MAPPED DATA')
df_existing_v1 = pd.read_excel(excel_v1, sheet_name='EXISTING BOX MAPPING')
df_new_v1 = pd.read_excel(excel_v1, sheet_name='NEW BOX CANDIDATES')
df_conflicts_v1 = pd.read_excel(excel_v1, sheet_name='MODEL CONFLICTS')

# Load Crosscheck Final Master for clean group titles
excel_cross = 'UZEE_TECH_SUPER_D_FINAL_MASTER_3AI_CROSSCHECK.xlsx'
df_fm = pd.read_excel(excel_cross, sheet_name='FINAL MASTER')
fm_dict = {r['Final Group ID']: r for idx, r in df_fm.iterrows()}

def extract_size(text):
    match = re.search(r'(\d+\.\d+)["\s]', text)
    if match:
        return f'{match.group(1)}"'
    return 'Unknown'

# --- DATASET PREPARATION ---

# Sheet 1: VERIFIED EXISTING BOXES
verified_boxes_rows = []
models_added_total = 0
all_models_in_verified = set()

for idx in range(106):
    row_final = df_final_v1.iloc[idx]
    row_exist = df_existing_v1.iloc[idx]
    
    bnum = row_final['BOX NUMBER']
    json_b = json_box_dict.get(bnum, {})
    
    matched_rg = row_exist['Matched New SD Group']
    rg_id = matched_rg.split(' ')[0] if pd.notna(matched_rg) else 'N/A'
    
    dsize = row_final['DISPLAY SIZE']
    title = row_final['TITLE']
    
    final_models_raw = str(row_final['COMPATIBLE MODELS'])
    final_models = [m.strip() for m in final_models_raw.split(',') if m.strip()]
    
    json_models = [m.strip() for m in json_b.get('compatibleModels', [])]
    new_added = set(final_models) - set(json_models)
    models_added_total += len(new_added)
    
    for m in final_models:
        all_models_in_verified.add(m)
        
    verified_boxes_rows.append({
        'Physical Box Number': bnum,
        'Research Group ID': rg_id,
        'Display Size': dsize if pd.notna(dsize) else 'Unknown',
        'Title': title if pd.notna(title) else 'Unknown',
        'Compatible Models': ", ".join(final_models),
        'Source': 'UZEE TECH Store Inventory + Mietubl Research Cross-Check',
        'Verification': 'Verified (Physical Box Active)'
    })

# Sheet 2: NEW GROUPS — BOX UNKNOWN
new_groups_rows = []
all_models_in_new_groups = set()

for idx in range(106, len(df_final_v1)):
    row_final = df_final_v1.iloc[idx]
    row_new = df_new_v1.iloc[idx - 106]
    
    rg_raw = row_new['Research SD Group']
    if '(' in str(rg_raw):
        rg_id = str(rg_raw).split('(')[0].strip()
    else:
        rg_id = str(rg_raw).strip()
        
    fm_info = fm_dict.get(rg_id, {})
    clean_title = fm_info.get('Super-D Glass Group')
    if not clean_title or pd.isna(clean_title):
        clean_title = str(row_final['TITLE']).replace('NEW GROUP: ', '').split(' (')[0].strip()
        
    dsize_orig = row_final['DISPLAY SIZE']
    if pd.notna(dsize_orig) and str(dsize_orig).strip() != 'Unknown':
        dsize = str(dsize_orig).strip()
    else:
        dsize = extract_size(str(clean_title))
        if dsize == 'Unknown':
            dsize = extract_size(str(row_final['COMPATIBLE MODELS']))
    
    models_raw = str(row_final['COMPATIBLE MODELS'])
    models = [m.strip() for m in models_raw.split(',') if m.strip()]
    for m in models:
        all_models_in_new_groups.add(m)
        
    evidence = row_new['Evidence'] if pd.notna(row_new['Evidence']) else 'Mietubl Research Group'
    
    new_groups_rows.append({
        'Research Group ID': rg_id,
        'Compatible Models': ", ".join(models),
        'Display Size': dsize,
        'Title': str(clean_title),
        'Sources': evidence,
        'Verification': 'BOX NUMBER REQUIRED',
        'Reason Box Number Is Unknown': 'No physical UZEE TECH box number assigned yet in store inventory / requires physical stock verification'
    })

# Sheet 3: MODEL CONFLICTS
conflicts_rows = []
for idx, r in df_conflicts_v1.iterrows():
    conflicts_rows.append({
        'Model': r['Phone Model'],
        'Existing Box': r['Existing Box'],
        'New Research Group': r['New Research Box/Group'],
        'Evidence': r['Conflict'],
        'Resolution Required': r['Recommended Action']
    })

# Sheet 4: COMPLETE MASTER
master_rows = []
for r in verified_boxes_rows:
    master_rows.append({
        'Physical Box Number': r['Physical Box Number'],
        'Research Group ID': r['Research Group ID'],
        'Display Size': r['Display Size'],
        'Title': r['Title'],
        'Compatible Models': r['Compatible Models'],
        'Source': r['Source'],
        'Verification': r['Verification']
    })

for r in new_groups_rows:
    master_rows.append({
        'Physical Box Number': 'UNKNOWN',
        'Research Group ID': r['Research Group ID'],
        'Display Size': r['Display Size'],
        'Title': r['Title'],
        'Compatible Models': r['Compatible Models'],
        'Source': r['Sources'],
        'Verification': r['Verification']
    })

# Statistics calculation
total_master_unique_models = all_models_in_verified.union(all_models_in_new_groups)

stats_rows = [
    {'Metric': 'Verified Physical Boxes (BOX 01–BOX 106)', 'Value': len(verified_boxes_rows), 'Category': 'Physical Inventory'},
    {'Metric': 'New Research Groups (Physical Box UNKNOWN)', 'Value': len(new_groups_rows), 'Category': 'Research Groups'},
    {'Metric': 'Total Master Super-D Groups', 'Value': len(master_rows), 'Category': 'Master Catalog'},
    {'Metric': 'Unique Models in Verified Physical Boxes', 'Value': len(all_models_in_verified), 'Category': 'Coverage'},
    {'Metric': 'Unique Models in Unknown Box Groups', 'Value': len(all_models_in_new_groups), 'Category': 'Coverage'},
    {'Metric': 'TOTAL UNIQUE MODELS ACROSS COMPLETE MASTER', 'Value': len(total_master_unique_models), 'Category': 'Coverage'},
    {'Metric': 'Total Compatible Model Additions to Physical Boxes', 'Value': models_added_total, 'Category': 'Model Expansion'},
    {'Metric': 'Identified Model Conflicts & Overlaps', 'Value': len(conflicts_rows), 'Category': 'Data Quality'}
]

# Convert to DataFrames
df_s1 = pd.DataFrame(verified_boxes_rows)
df_s2 = pd.DataFrame(new_groups_rows)
df_s3 = pd.DataFrame(conflicts_rows)
df_s4 = pd.DataFrame(master_rows)
df_s5 = pd.DataFrame(stats_rows)

print(f"Sheet 1 (VERIFIED EXISTING BOXES): {len(df_s1)} rows")
print(f"Sheet 2 (NEW GROUPS — BOX UNKNOWN): {len(df_s2)} rows")
print(f"Sheet 3 (MODEL CONFLICTS): {len(df_s3)} rows")
print(f"Sheet 4 (COMPLETE MASTER): {len(df_s4)} rows")
print(f"Sheet 5 (DATA STATISTICS): {len(df_s5)} rows")

# --- EXCEL STYLING WITH OPENPYXL ---

output_file = 'UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx'

wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

# Colors & Fonts
navy_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

verified_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
verified_font = Font(name="Calibri", size=10, color="166534", bold=True)

unknown_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
unknown_font = Font(name="Calibri", size=10, color="92400E", bold=True)

thin_border_side = Side(border_style="thin", color="CBD5E1")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

data_font = Font(name="Calibri", size=10, color="0F172A")

sheets_data = [
    ("VERIFIED EXISTING BOXES", df_s1),
    ("NEW GROUPS — BOX UNKNOWN", df_s2),
    ("MODEL CONFLICTS", df_s3),
    ("COMPLETE MASTER", df_s4),
    ("DATA STATISTICS", df_s5)
]

for sheet_title, df_sheet in sheets_data:
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    headers = list(df_sheet.columns)
    ws.append(headers)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 28
    
    for row_idx, row_data in enumerate(df_sheet.values, start=2):
        ws.append(list(row_data))
        ws.row_dimensions[row_idx].height = 22
        
        row_bg = zebra_fill if row_idx % 2 == 0 else white_fill
        
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = row_bg
            cell.border = thin_border
            
            header_name = headers[col_idx - 1]
            
            if header_name in ['Physical Box Number', 'Research Group ID', 'Display Size', 'Verification', 'Metric', 'Value', 'Category']:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
            if header_name in ['Verification', 'Verification Status']:
                val_str = str(val).upper()
                if 'VERIFIED' in val_str:
                    cell.fill = verified_fill
                    cell.font = verified_font
                elif 'REQUIRED' in val_str or 'UNKNOWN' in val_str:
                    cell.fill = unknown_fill
                    cell.font = unknown_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df_sheet) + 1}"

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        header_text = str(col[0].value)
        
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
                
        if header_text in ['Compatible Models', 'Sources', 'Reason Box Number Is Unknown', 'Evidence', 'Resolution Required', 'Source']:
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 25), 65)
        elif header_text in ['Title', 'New Research Group', 'Metric']:
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 20), 40)
        else:
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

wb.save(output_file)
print(f"\n🎉 Successfully saved revised workbook '{output_file}'!")
