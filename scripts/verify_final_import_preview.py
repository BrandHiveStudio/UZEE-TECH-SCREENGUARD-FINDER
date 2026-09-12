import openpyxl
import pandas as pd
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_name = 'UZEE_TECH_SUPER_D_FINAL_IMPORT_PREVIEW.xlsx'
wb = openpyxl.load_workbook(file_name, data_only=True)

print("Final Import Preview Workbook Sheet Names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    df = pd.read_excel(file_name, sheet_name=sname)
    print(f"\n--- Sheet '{sname}' ---")
    print(f"  Rows (excl header): {ws.max_row - 1}, Cols: {ws.max_column}")
    print("  Columns:", df.columns.tolist())
    print("  First 2 rows:")
    print(df.head(2).to_string())
    print("  Freeze panes:", ws.freeze_panes)
    print("  AutoFilter ref:", ws.auto_filter.ref)

print("\n-----------------------------------------------------")
print("DATA INTEGRITY VERIFICATION:")

df_s1 = pd.read_excel(file_name, sheet_name='CURRENT DATA')
df_s2 = pd.read_excel(file_name, sheet_name='NEW DATA')
df_s3 = pd.read_excel(file_name, sheet_name='ADDED GROUPS')
df_s4 = pd.read_excel(file_name, sheet_name='ADDED MODEL RELATIONSHIPS')
df_s5 = pd.read_excel(file_name, sheet_name='CHANGED RECORDS')
df_s6 = pd.read_excel(file_name, sheet_name='POTENTIAL CONFLICTS')
df_s7 = pd.read_excel(file_name, sheet_name='VALIDATION RESULTS')
df_s8 = pd.read_excel(file_name, sheet_name='FINAL COUNTS')

print(f"Current Data Rows: {len(df_s1)}")
print(f"New Data Rows: {len(df_s2)}")
print(f"Added Groups Rows: {len(df_s3)}")
print(f"Added Model Relationships Rows: {len(df_s4)}")
print(f"Changed Records Rows: {len(df_s5)}")
print(f"Potential Conflicts Rows: {len(df_s6)}")
print(f"Validation Checks Passed: {len(df_s7[df_s7['Status'] == 'PASSED'])} / {len(df_s7)}")
print(f"Final Counts Rows: {len(df_s8)}")

assert len(df_s1) == 106, f"Expected 106 production records, got {len(df_s1)}"
assert len(df_s2) == 263, f"Expected 263 new data records, got {len(df_s2)}"
assert len(df_s7[df_s7['Status'] == 'PASSED']) == 8, "Expected 8 passed validation checks"

print("\n✅ FINAL IMPORT PREVIEW VERIFICATION PASSED PERFECTLY!")
