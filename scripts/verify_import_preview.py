import openpyxl
import pandas as pd
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_name = 'UZEE_TECH_SUPER_D_IMPORT_PREVIEW.xlsx'
wb = openpyxl.load_workbook(file_name, data_only=True)

print("Import Preview Workbook Sheet Names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    df = pd.read_excel(file_name, sheet_name=sname)
    print(f"\n--- Sheet '{sname}' ---")
    print(f"  Rows (excl header): {ws.max_row - 1}, Cols: {ws.max_column}")
    print("  Columns:", df.columns.tolist())
    print("  First 2 rows:")
    print(df.head(2).to_string())
    print("  Freeze panes:", ws.freeze_panes)
    print("  AutoFilter ref:", ws.auto_filter.ref)

print("\n-----------------------------------------------------")
print("DATA INTEGRITY VERIFICATION:")

df_s1 = pd.read_excel(file_name, sheet_name='CURRENT PRODUCTION DATA')
df_s2 = pd.read_excel(file_name, sheet_name='AUTHORITATIVE IMPORT DATA')
df_s3 = pd.read_excel(file_name, sheet_name='ADDED CHANGED REMOVED COMP')
df_s4 = pd.read_excel(file_name, sheet_name='MODEL RELATIONSHIP COMPARISON')
df_s5 = pd.read_excel(file_name, sheet_name='VALIDATION RESULTS')

print(f"Current Production Boxes: {len(df_s1)}")
print(f"Authoritative Import Boxes: {len(df_s2)}")
print(f"Diff Relationships Rows: {len(df_s3)}")
print(f"Model Relationship Comparison Rows: {len(df_s4)}")
print(f"Validation Checks Passed: {len(df_s5[df_s5['Status'] == 'PASSED'])} / {len(df_s5)}")

assert len(df_s1) == 106, f"Expected 106 production boxes, got {len(df_s1)}"
assert len(df_s2) == 106, f"Expected 106 import boxes, got {len(df_s2)}"
assert len(df_s5[df_s5['Status'] == 'PASSED']) == 8, "Expected 8 passed validation checks"

print("\n✅ IMPORT PREVIEW VERIFICATION PASSED PERFECTLY!")
