import openpyxl
import pandas as pd
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx', data_only=True)

print("Workbook sheet names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    df = pd.read_excel('UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx', sheet_name=sname)
    print(f"\n--- Sheet '{sname}' ---")
    print(f"  Rows (excl header): {ws.max_row - 1}, Cols: {ws.max_column}")
    print("  Columns:", df.columns.tolist())
    print("  First 2 rows:")
    print(df.head(2).to_string())
    print("  Freeze panes:", ws.freeze_panes)
    print("  AutoFilter ref:", ws.auto_filter.ref)

print("\n-----------------------------------------------------")
print("DATA INTEGRITY VERIFICATION:")

df_s1 = pd.read_excel('UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx', sheet_name='VERIFIED EXISTING BOXES')
df_s2 = pd.read_excel('UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx', sheet_name='NEW GROUPS — BOX UNKNOWN')
df_s3 = pd.read_excel('UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx', sheet_name='MODEL CONFLICTS')
df_s4 = pd.read_excel('UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx', sheet_name='COMPLETE MASTER')
df_s5 = pd.read_excel('UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx', sheet_name='DATA STATISTICS')

# Verify no BOX 107 in Sheet 2
unknown_boxes_in_s2 = df_s2['Verification'].unique()
print(f"Sheet 2 Verification values: {unknown_boxes_in_s2}")

master_box_numbers = df_s4['Physical Box Number'].unique()
print(f"Sheet 4 Box Numbers sample: {master_box_numbers[:10]} ... {master_box_numbers[-5:]}")

# Verify statistics
verified_count = len(df_s1)
unknown_groups_count = len(df_s2)
model_conflicts_count = len(df_s3)

models_in_s1 = set()
for mlist in df_s1['Compatible Models']:
    for m in str(mlist).split(','):
        models_in_s1.add(m.strip())

models_in_s2 = set()
for mlist in df_s2['Compatible Models']:
    for m in str(mlist).split(','):
        models_in_s2.add(m.strip())

total_unique_models = models_in_s1.union(models_in_s2)

print("\nEXACT REPORTED METRICS:")
print(f"- Verified physical boxes: {verified_count}")
print(f"- New groups with unknown box numbers: {unknown_groups_count}")
print(f"- Total unique models: {len(total_unique_models)}")
print(f"- Model conflicts: {model_conflicts_count}")
print(f"- Models added to existing boxes: {df_s5[df_s5['Metric'] == 'Total Compatible Model Additions to Physical Boxes']['Value'].values[0]}")
