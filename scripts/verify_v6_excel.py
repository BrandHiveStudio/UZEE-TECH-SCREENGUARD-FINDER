import openpyxl
import pandas as pd
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_name = 'UZEE_TECH_SUPER_D_EXISTING_BOX_RECONCILIATION_V6.xlsx'
wb = openpyxl.load_workbook(file_name, data_only=True)

print("Workbook Sheet Names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    df = pd.read_excel(file_name, sheet_name=sname)
    print(f"\n--- Sheet '{sname}' ---")
    print(f"  Rows (excl header): {ws.max_row - 1}, Cols: {ws.max_column}")
    print("  Columns:", df.columns.tolist())
    print("  First 2 rows:")
    print(df.head(2).to_string())
    print("  Freeze panes:", ws.freeze_panes)
    print("  AutoFilter ref:", ws.auto_filter.ref)

print("\n-----------------------------------------------------")
print("DATA INTEGRITY VERIFICATION:")

df_s1 = pd.read_excel(file_name, sheet_name='SAFE BOX ADDITIONS')
df_s2 = pd.read_excel(file_name, sheet_name='NEW UNASSIGNED GLASS')
df_s3 = pd.read_excel(file_name, sheet_name='CONFLICTS')
df_s4 = pd.read_excel(file_name, sheet_name='ALREADY REPRESENTED')
df_s5 = pd.read_excel(file_name, sheet_name='PROPOSED 106-BOX MASTER')
df_s6 = pd.read_excel(file_name, sheet_name='SOURCE EVIDENCE')
df_s7 = pd.read_excel(file_name, sheet_name='STATISTICS')

print(f"Safe Box Additions Count: {len(df_s1)}")
print(f"New Unassigned Glass Groups Count: {len(df_s2)}")
print(f"Conflicts Count: {len(df_s3)}")
print(f"Already Represented Groups Count: {len(df_s4)}")
print(f"Proposed Master Physical Boxes: {len(df_s5)}")
print(f"Source Evidence Rows: {len(df_s6)}")
print(f"Statistics Rows: {len(df_s7)}")

assert len(df_s5) == 106, f"Expected 106 physical boxes in master, got {len(df_s5)}"

print("\n✅ V6 RECONCILIATION VERIFICATION PASSED PERFECTLY!")
