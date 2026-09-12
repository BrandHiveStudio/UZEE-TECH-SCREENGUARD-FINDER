import openpyxl
import pandas as pd
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_name = 'UZEE_TECH_SUPER_D_BOX_MAPPING_V3_CLEAN.xlsx'
wb = openpyxl.load_workbook(file_name, data_only=True)

print("Workbook Sheet Names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    df = pd.read_excel(file_name, sheet_name=sname)
    print(f"\n--- Sheet '{sname}' ---")
    print(f"  Rows (excl header): {ws.max_row - 1}, Cols: {ws.max_column}")
    print("  Columns:", df.columns.tolist())
    print("  First 2 rows:")
    print(df.head(2).to_string())
    print("  Freeze panes:", ws.freeze_panes)
    print("  AutoFilter ref:", ws.auto_filter.ref)

print("\n-----------------------------------------------------")
print("DATA INTEGRITY VERIFICATION:")

df_s1 = pd.read_excel(file_name, sheet_name='FINAL CLEAN BOX DATA')
df_s2 = pd.read_excel(file_name, sheet_name='MODEL LOOKUP')
df_s3 = pd.read_excel(file_name, sheet_name='ALREADY REPRESENTED')
df_s4 = pd.read_excel(file_name, sheet_name='NEW COMPATIBILITY')
df_s5 = pd.read_excel(file_name, sheet_name='MODEL MULTI-BOX REVIEW')
df_s6 = pd.read_excel(file_name, sheet_name='CONFLICTS')
df_s7 = pd.read_excel(file_name, sheet_name='NORMALIZATION LOG')
df_s8 = pd.read_excel(file_name, sheet_name='DATA STATISTICS')

# Metrics check
physical_boxes = len(df_s1)
unique_canonical_models_in_boxes = df_s2['Canonical Model'].nunique()
total_relationships = len(df_s2)
already_rep = len(df_s3)
new_compat = len(df_s4)
multi_box_models = len(df_s5)
conflicts_count = len(df_s6)
dup_removed = df_s8[df_s8['Metric'] == 'Duplicate Names Removed Within Boxes']['Value'].values[0]

print(f"Physical Boxes: {physical_boxes}")
print(f"Canonical Unique Models in Physical Boxes: {unique_canonical_models_in_boxes}")
print(f"Total Model-to-Box Relationships: {total_relationships}")
print(f"Unknown Groups Total: {already_rep + new_compat}")
print(f"Already Represented Groups: {already_rep}")
print(f"New Compatibility Groups: {new_compat}")
print(f"Conflicts: {conflicts_count}")
print(f"Multi-Box Models: {multi_box_models}")
print(f"Duplicate Names Removed Within Boxes: {dup_removed}")
