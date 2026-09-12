import pandas as pd
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')

# 1. Load V2 Excel file
excel_v2 = 'UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS_V2.xlsx'
df_s1_v2 = pd.read_excel(excel_v2, sheet_name='VERIFIED EXISTING BOXES')
df_s2_v2 = pd.read_excel(excel_v2, sheet_name='NEW GROUPS — BOX UNKNOWN')
df_s3_v2 = pd.read_excel(excel_v2, sheet_name='MODEL CONFLICTS')

# Load Crosscheck Final Master for clean group titles
excel_cross = 'UZEE_TECH_SUPER_D_FINAL_MASTER_3AI_CROSSCHECK.xlsx'
df_fm = pd.read_excel(excel_cross, sheet_name='FINAL MASTER')
fm_dict = {r['Final Group ID']: r for idx, r in df_fm.iterrows()}

# Normalization Function
def normalize_model(raw):
    m = str(raw).strip()
    
    # Clean factory text annotations attached to model names
    m = re.sub(r'\s*\d*胶$', '', m)
    m = re.sub(r'\s*玻璃[\d\.]+MM$', '', m, flags=re.IGNORECASE)
    m = re.sub(r'\s*THICK\s+GLUE$', '', m, flags=re.IGNORECASE)

    # Standardize brand prefixes & double brand prefixes
    m = re.sub(r'^(Samsung|SAMSUNG)\s+SAM\s+', 'Samsung ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(Redmi|REDMI)\s+RM\s+', 'Redmi ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(iPhone|IPHONE)\s+IP\s+', 'iPhone ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(OPPO|Oppo)\s+OP\s+', 'OPPO ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(Vivo|VIVO)\s+VO\s+', 'Vivo ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(Realme|REALME)\s+REAL\s+', 'Realme ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(POCO|Poco)\s+POC\s+', 'POCO ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(OnePlus|ONEPLUS)\s+1\+\s+', 'OnePlus ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(Xiaomi|XIAOMI)\s+XM\s+', 'Xiaomi ', m, flags=re.IGNORECASE)

    m = re.sub(r'^(IP|IPHONE)\s+', 'iPhone ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(SAM|SAMSUNG)\s+', 'Samsung ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(RM|REDMI)\s+', 'Redmi ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(OP|OPPO)\s+', 'OPPO ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(VO|VIVO)\s+', 'Vivo ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(REAL|REALME)\s+', 'Realme ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(POC|POCO)\s+', 'POCO ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(1\+|ONEPLUS)\s+', 'OnePlus ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(XM|XIAOMI)\s+', 'Xiaomi ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(HONOR)\s+', 'Honor ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(HUAWEI)\s+', 'Huawei ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(MOTO|MOTOROLA)\s+', 'Motorola ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(ITEL)\s+', 'Itel ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(INFINIX)\s+', 'Infinix ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(TECNO)\s+', 'Tecno ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(GOOGLE)\s+', 'Google ', m, flags=re.IGNORECASE)
    m = re.sub(r'^(NOKIA)\s+', 'Nokia ', m, flags=re.IGNORECASE)

    if re.match(r'^(HOT|SMART)\s+\d+', m, re.IGNORECASE):
        m = 'Infinix ' + m
    elif re.match(r'^(SPARK|POP)\s+\d+', m, re.IGNORECASE):
        m = 'Tecno ' + m

    def fix_casing(match):
        word = match.group(0).upper()
        mapping = {
            'PLUS': 'Plus', 'PRO': 'Pro', 'MINI': 'Mini', 'MAX': 'Max', 'LITE': 'Lite',
            'ULTRA': 'Ultra', 'POWER': 'Power', 'PRIME': 'Prime', 'PLAY': 'Play',
            'NOTE': 'Note', 'NEO': 'Neo', 'ZOOM': 'Zoom', 'MAGIC': 'Magic', 'PIXEL': 'Pixel',
            'SMART': 'Smart', 'SPARK': 'Spark', 'HOT': 'Hot', 'POP': 'Pop', 'ENJOY': 'Enjoy',
            'YOUTH': 'Youth', 'FE': 'FE', 'GT': 'GT', 'SE': 'SE', 'INDIA': 'India',
            'CHINA': 'China', 'GLOBAL': 'Global'
        }
        return mapping.get(word, match.group(0).capitalize())

    keywords = r'\b(PLUS|PRO|MINI|MAX|LITE|ULTRA|POWER|PRIME|PLAY|NOTE|NEO|ZOOM|MAGIC|PIXEL|SMART|SPARK|HOT|POP|ENJOY|YOUTH|FE|GT|SE|INDIA|CHINA|GLOBAL)\b'
    m = re.sub(keywords, fix_casing, m, flags=re.IGNORECASE)
    m = re.sub(r'\s+', ' ', m).strip()
    return m

def get_brand(model_str):
    tokens = model_str.split(' ')
    return tokens[0] if tokens else 'Unknown'

# --- 1. SHEET 1: FINAL CLEAN BOX DATA ---
clean_box_rows = []
box_to_models_map = {}
normalization_log = []
duplicates_removed_count = 0
total_relationships_s1 = 0

for idx, r in df_s1_v2.iterrows():
    bnum = r['Physical Box Number']
    rg_id = r['Research Group ID']
    dsize = r['Display Size']
    title = r['Title']
    source = r['Source']
    verif = r['Verification']
    
    raw_models_list = [m.strip() for m in str(r['Compatible Models']).split(',') if m.strip()]
    
    canon_models_list = []
    seen_in_box = set()
    
    for rm in raw_models_list:
        cm = normalize_model(rm)
        if rm != cm:
            normalization_log.append({
                'Original Model': rm,
                'Canonical Model': cm,
                'Reason': 'Standardized brand prefix/casing & removed factory annotations'
            })
            
        if cm in seen_in_box:
            duplicates_removed_count += 1
        else:
            seen_in_box.add(cm)
            canon_models_list.append(cm)
            
    box_to_models_map[bnum] = (rg_id, seen_in_box)
    total_relationships_s1 += len(canon_models_list)
    
    clean_box_rows.append({
        'Physical Box Number': bnum,
        'Canonical Compatible Models': ", ".join(canon_models_list),
        'Display Size': dsize,
        'Title': title,
        'Source': source,
        'Verification': verif,
        'Notes': f"{len(canon_models_list)} canonical models mapped; duplicate naming variations merged"
    })

# --- 2. SHEET 2: MODEL LOOKUP ---
model_lookup_rows = []
model_to_boxes_map = {}

for r in clean_box_rows:
    bnum = r['Physical Box Number']
    rg_id = box_to_models_map[bnum][0]
    c_models = [m.strip() for m in r['Canonical Compatible Models'].split(',') if m.strip()]
    
    for cm in c_models:
        brand = get_brand(cm)
        model_lookup_rows.append({
            'Canonical Model': cm,
            'Brand': brand,
            'Physical Box Number': bnum,
            'Verification': 'Verified (Physical Box Active)',
            'Source': 'UZEE TECH Store Inventory + Mietubl Research Cross-Check',
            'Research Group ID': rg_id
        })
        model_to_boxes_map.setdefault(cm, []).append(bnum)

# --- 3. RECONCILE SHEET 3 (ALREADY REPRESENTED) vs SHEET 4 (NEW COMPATIBILITY) ---
already_represented_rows = []
new_compatibility_rows = []
box_models_union = set(model_to_boxes_map.keys())

for idx, r in df_s2_v2.iterrows():
    rg_id = r['Research Group ID']
    title = r['Title']
    dsize = r['Display Size']
    evidence = r['Sources']
    
    raw_models = [m.strip() for m in str(r['Compatible Models']).split(',') if m.strip()]
    canon_models = list(dict.fromkeys([normalize_model(m) for m in raw_models]))
    
    for rm in raw_models:
        cm = normalize_model(rm)
        if rm != cm:
            normalization_log.append({
                'Original Model': rm,
                'Canonical Model': cm,
                'Reason': 'Standardized brand prefix/casing & removed factory annotations'
            })

    matching_boxes = {}
    for cm in canon_models:
        if cm in model_to_boxes_map:
            for b in model_to_boxes_map[cm]:
                matching_boxes.setdefault(b, []).append(cm)
                
    if matching_boxes:
        best_box, matched_models = max(matching_boxes.items(), key=lambda x: len(x[1]))
        overlap_pct = len(matched_models) / len(canon_models)
        
        if overlap_pct >= 0.5 or len(matched_models) == len(canon_models):
            already_represented_rows.append({
                'Unknown Group': f"{rg_id} ({title})",
                'Existing Box': best_box,
                'Matching Models': ", ".join(matched_models),
                'Reason': f"High compatibility overlap ({len(matched_models)}/{len(canon_models)} models match existing {best_box} inventory)",
                'Action': f"No new physical box required; map directly to existing physical {best_box}"
            })
        else:
            new_models = [cm for cm in canon_models if cm not in box_models_union]
            new_compatibility_rows.append({
                'Research Group': f"{rg_id} ({title})",
                'New Models': ", ".join(new_models) if new_models else ", ".join(canon_models),
                'Evidence': evidence,
                'Physical Box Number': 'UNKNOWN',
                'Status': 'BOX NUMBER REQUIRED'
            })
    else:
        new_compatibility_rows.append({
            'Research Group': f"{rg_id} ({title})",
            'New Models': ", ".join(canon_models),
            'Evidence': evidence,
            'Physical Box Number': 'UNKNOWN',
            'Status': 'BOX NUMBER REQUIRED'
        })

# --- 4. SHEET 5: MODEL MULTI-BOX REVIEW ---
multi_box_models_s1 = {m: boxes for m, boxes in model_to_boxes_map.items() if len(boxes) > 1}
multi_box_review_rows = []

for cm, boxes in multi_box_models_s1.items():
    multi_box_review_rows.append({
        'Canonical Model': cm,
        'Box 1': boxes[0],
        'Box 2': boxes[1] if len(boxes) > 1 else 'N/A',
        'Evidence': f"Model mapped to multiple physical boxes in inventory ({', '.join(boxes)})",
        'Likely Legitimate?': 'Yes (Multi-Box Stocking / Dual Fit)',
        'Resolution': 'Retain multi-box mapping in database; verify physical glass dimension differences'
    })

# --- 5. SHEET 6: CONFLICTS ---
conflicts_rows = []
for idx, r in df_s3_v2.iterrows():
    m_norm = normalize_model(r['Model'])
    conflicts_rows.append({
        'Model': m_norm,
        'Existing Box': r['Existing Box'],
        'Research Group': r['New Research Group'],
        'Conflict': r['Evidence'],
        'Evidence': r['Evidence'],
        'Required Action': r['Resolution Required']
    })

# --- 6. SHEET 7: NORMALIZATION LOG ---
unique_norm_log = []
seen_log_keys = set()
for log_entry in normalization_log:
    key = (log_entry['Original Model'], log_entry['Canonical Model'])
    if key not in seen_log_keys:
        seen_log_keys.add(key)
        unique_norm_log.append(log_entry)

# --- 7. SHEET 8: DATA STATISTICS ---
stats_rows = [
    {'Metric': 'Physical Boxes (BOX 01–BOX 106)', 'Value': len(clean_box_rows), 'Category': 'Physical Inventory'},
    {'Metric': 'Canonical Unique Models in Physical Boxes', 'Value': len(model_to_boxes_map), 'Category': 'Coverage'},
    {'Metric': 'Total Model-to-Box Relationships', 'Value': total_relationships_s1, 'Category': 'Relationships'},
    {'Metric': 'Unknown Research Groups (Total)', 'Value': len(df_s2_v2), 'Category': 'Research Groups'},
    {'Metric': 'Already Represented Unknown Groups', 'Value': len(already_represented_rows), 'Category': 'Reconciliation'},
    {'Metric': 'New Compatibility Unknown Groups', 'Value': len(new_compatibility_rows), 'Category': 'Reconciliation'},
    {'Metric': 'Identified Conflicts & Overlaps', 'Value': len(conflicts_rows), 'Category': 'Data Quality'},
    {'Metric': 'Multi-Box Models in Physical Inventory', 'Value': len(multi_box_models_s1), 'Category': 'Multi-Box Relationships'},
    {'Metric': 'Duplicate Names Removed Within Boxes', 'Value': duplicates_removed_count, 'Category': 'Normalization'}
]

# Create DataFrames
df_s1 = pd.DataFrame(clean_box_rows)
df_s2 = pd.DataFrame(model_lookup_rows)
df_s3 = pd.DataFrame(already_represented_rows)
df_s4 = pd.DataFrame(new_compatibility_rows)
df_s5 = pd.DataFrame(multi_box_review_rows)
df_s6 = pd.DataFrame(conflicts_rows)
df_s7 = pd.DataFrame(unique_norm_log)
df_s8 = pd.DataFrame(stats_rows)

print(f"Sheet 1 (FINAL CLEAN BOX DATA): {len(df_s1)} rows")
print(f"Sheet 2 (MODEL LOOKUP): {len(df_s2)} rows")
print(f"Sheet 3 (ALREADY REPRESENTED): {len(df_s3)} rows")
print(f"Sheet 4 (NEW COMPATIBILITY): {len(df_s4)} rows")
print(f"Sheet 5 (MODEL MULTI-BOX REVIEW): {len(df_s5)} rows")
print(f"Sheet 6 (CONFLICTS): {len(df_s6)} rows")
print(f"Sheet 7 (NORMALIZATION LOG): {len(df_s7)} rows")
print(f"Sheet 8 (DATA STATISTICS): {len(df_s8)} rows")

# --- EXCEL STYLING WITH OPENPYXL ---

output_file = 'UZEE_TECH_SUPER_D_BOX_MAPPING_V3_CLEAN.xlsx'

wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

navy_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

verified_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
verified_font = Font(name="Calibri", size=10, color="166534", bold=True)

unknown_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
unknown_font = Font(name="Calibri", size=10, color="92400E", bold=True)

thin_border_side = Side(border_style="thin", color="CBD5E1")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

data_font = Font(name="Calibri", size=10, color="0F172A")

sheets_data = [
    ("FINAL CLEAN BOX DATA", df_s1),
    ("MODEL LOOKUP", df_s2),
    ("ALREADY REPRESENTED", df_s3),
    ("NEW COMPATIBILITY", df_s4),
    ("MODEL MULTI-BOX REVIEW", df_s5),
    ("CONFLICTS", df_s6),
    ("NORMALIZATION LOG", df_s7),
    ("DATA STATISTICS", df_s8)
]

for sheet_title, df_sheet in sheets_data:
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    headers = list(df_sheet.columns)
    ws.append(headers)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 28
    
    for row_idx, row_data in enumerate(df_sheet.values, start=2):
        ws.append(list(row_data))
        ws.row_dimensions[row_idx].height = 22
        
        row_bg = zebra_fill if row_idx % 2 == 0 else white_fill
        
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = row_bg
            cell.border = thin_border
            
            header_name = headers[col_idx - 1]
            
            if header_name in ['Physical Box Number', 'Research Group ID', 'Display Size', 'Verification', 'Status', 'Metric', 'Value', 'Category', 'Brand', 'Box 1', 'Box 2']:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
            if header_name in ['Verification', 'Verification Status', 'Status']:
                val_str = str(val).upper()
                if 'VERIFIED' in val_str:
                    cell.fill = verified_fill
                    cell.font = verified_font
                elif 'REQUIRED' in val_str or 'UNKNOWN' in val_str:
                    cell.fill = unknown_fill
                    cell.font = unknown_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df_sheet) + 1}"

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        header_text = str(col[0].value)
        
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
                
        if header_text in ['Canonical Compatible Models', 'Sources', 'Evidence', 'Resolution Required', 'Source', 'Reason', 'Action', 'Matching Models', 'New Models', 'Notes']:
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 25), 65)
        elif header_text in ['Title', 'Canonical Model', 'Original Model', 'Unknown Group', 'Research Group', 'Metric']:
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 20), 40)
        else:
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

wb.save(output_file)
print(f"\n🎉 Successfully saved revised workbook '{output_file}' with 8 professional sheets!")
