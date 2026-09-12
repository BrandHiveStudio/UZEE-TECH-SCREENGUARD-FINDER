import json
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('UZEE_TECH_SUPER_D_MASTER_263_WITH_DISPLAY_SIZE_QA.xlsx', data_only=True)
ws = wb['MASTER_263']

excel_ds = {}
for r in range(2, ws.max_row + 1):
    gid = ws.cell(row=r, column=1).value
    dsize = ws.cell(row=r, column=5).value
    if gid:
        excel_ds[str(gid).strip()] = str(dsize).strip() if dsize else 'Unknown'

print(f"Loaded {len(excel_ds)} display sizes from Excel sheet MASTER_263.")

with open('src/data/screenguards.json', 'r', encoding='utf-8') as f:
    json_data = json.load(f)

boxes = json_data['boxes']
updated_count = 0
unknown_count = 0

for b in boxes:
    gid = b['id']
    if gid in excel_ds:
        new_ds = excel_ds[gid]
        if b['displaySize'] != new_ds:
            b['displaySize'] = new_ds
            updated_count += 1
        if new_ds == 'Unknown':
            unknown_count += 1

with open('src/data/screenguards.json', 'w', encoding='utf-8') as f:
    json.dump(json_data, f, indent=2, ensure_ascii=False)

print(f"Updated {updated_count} display sizes in src/data/screenguards.json.")
print(f"Total groups: {len(boxes)}")
print(f"Groups with known display size: {len(boxes) - unknown_count}")
print(f"Groups with Unknown display size: {unknown_count}")
