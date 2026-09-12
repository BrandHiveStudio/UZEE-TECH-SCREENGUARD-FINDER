import openpyxl
import pandas as pd
import sys

sys.stdout.reconfigure(encoding='utf-8')

master_file = 'UZEE_TECH_SUPER_D_COMPLETE_EDITABLE_MASTER.xlsx'
wb = openpyxl.load_workbook(master_file, data_only=True)

print("Sheet Names in UZEE_TECH_SUPER_D_COMPLETE_EDITABLE_MASTER.xlsx:")
print(wb.sheetnames)

for sname in wb.sheetnames:
    df = pd.read_excel(master_file, sheet_name=sname)
    print(f"\n--- Sheet '{sname}' ---")
    print(f"  Rows: {len(df)}, Columns ({len(df.columns)}): {df.columns.tolist()}")
    print("  First 2 rows:")
    print(df.head(2).to_string())
