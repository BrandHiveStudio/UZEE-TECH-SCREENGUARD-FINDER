import openpyxl

for fname in ['UZEE_TECH_SUPER_D_BOX_MAPPING_ANALYSIS.xlsx', 'UZEE_TECH_SUPER_D_FINAL_MASTER_3AI_CROSSCHECK.xlsx']:
    try:
        wb = openpyxl.load_workbook(fname, data_only=True)
        print(f'=== {fname} ===')
        print('Sheets:', wb.sheetnames)
        for sname in wb.sheetnames:
            ws = wb[sname]
            headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
            print(f'  Sheet "{sname}": {ws.max_row} rows, {ws.max_column} cols')
            print(f'    Headers: {headers[:10]}')
    except Exception as e:
        print(f'Error loading {fname}: {e}')
